#!/usr/bin/env python3
"""Build the compact, public ComReg mobile site catalogue used by the map.

Input workbooks are the non-confidential 2026 Q1 mobile licence schedules from
ComReg. Keep the source workbooks out of the published site; only the derived
coordinates, licensed bands, service flags, and maximum EIRP are included.
Requires openpyxl. Run with: python scripts/build-mobile-sites.py <xlsx-dir> <output-json>
"""

from __future__ import annotations

import json
import math
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


FILES = {
    "eir_mbsa1_2026.xlsx": "Eir",
    "eir_mbsa2_2026.xlsx": "Eir",
    "eir_36ghz_2026.xlsx": "Eir",
    "three_mbsa1_a_2026.xlsx": "Three",
    "three_mbsa1_b_2026.xlsx": "Three",
    "three_mbsa2_2026.xlsx": "Three",
    "three_36ghz_2026.xlsx": "Three",
    "vodafone_mbsa1_2026.xlsx": "Vodafone",
    "vodafone_mbsa2_2026.xlsx": "Vodafone",
    "vodafone_36ghz_2026.xlsx": "Vodafone",
}
NETWORKS = {"Eir": 0, "Three": 1, "Vodafone": 2}
SOURCE_URL = "https://www.comreg.ie/industry/radio-spectrum/licensing/search-licence-type/mobile-licences-2/"

# TM75 / Irish Grid, EPSG:29903, followed by the published 7-parameter
# transformation to WGS84.
A = 6377340.189
INV_F = 299.3249646
F = 1 / INV_F
E2 = 2 * F - F * F
EP2 = E2 / (1 - E2)
K0 = 1.000035
PHI0 = math.radians(53.5)
LAM0 = math.radians(-8.0)


def meridional_arc(phi: float) -> float:
    return A * (
        (1 - E2 / 4 - 3 * E2**2 / 64 - 5 * E2**3 / 256) * phi
        - (3 * E2 / 8 + 3 * E2**2 / 32 + 45 * E2**3 / 1024) * math.sin(2 * phi)
        + (15 * E2**2 / 256 + 45 * E2**3 / 1024) * math.sin(4 * phi)
        - (35 * E2**3 / 3072) * math.sin(6 * phi)
    )


M0 = meridional_arc(PHI0)


def irish_grid_to_wgs84(easting: float, northing: float) -> tuple[float, float]:
    x = (easting - 200000) / K0
    m = M0 + (northing - 250000) / K0
    mu = m / (A * (1 - E2 / 4 - 3 * E2**2 / 64 - 5 * E2**3 / 256))
    e1 = (1 - math.sqrt(1 - E2)) / (1 + math.sqrt(1 - E2))
    phi1 = (
        mu
        + (3 * e1 / 2 - 27 * e1**3 / 32) * math.sin(2 * mu)
        + (21 * e1**2 / 16 - 55 * e1**4 / 32) * math.sin(4 * mu)
        + (151 * e1**3 / 96) * math.sin(6 * mu)
        + (1097 * e1**4 / 512) * math.sin(8 * mu)
    )
    sin_phi, cos_phi = math.sin(phi1), math.cos(phi1)
    tan_phi = math.tan(phi1)
    n1 = A / math.sqrt(1 - E2 * sin_phi**2)
    r1 = A * (1 - E2) / (1 - E2 * sin_phi**2) ** 1.5
    t1 = tan_phi**2
    c1 = EP2 * cos_phi**2
    d = x / n1
    lat = phi1 - (n1 * tan_phi / r1) * (
        d**2 / 2
        - (5 + 3 * t1 + 10 * c1 - 4 * c1**2 - 9 * EP2) * d**4 / 24
        + (61 + 90 * t1 + 298 * c1 + 45 * t1**2 - 252 * EP2 - 3 * c1**2) * d**6 / 720
    )
    lon = LAM0 + (
        d
        - (1 + 2 * t1 + c1) * d**3 / 6
        + (5 - 2 * c1 + 28 * t1 - 3 * c1**2 + 8 * EP2 + 24 * t1**2) * d**5 / 120
    ) / cos_phi

    # Source ellipsoid geodetic -> cartesian.
    xx = n1 * math.cos(lat) * math.cos(lon)
    yy = n1 * math.cos(lat) * math.sin(lon)
    zz = n1 * (1 - E2) * math.sin(lat)
    tx, ty, tz = 482.5, -130.6, 564.6
    rx, ry, rz = [math.radians(v / 3600) for v in (-1.042, -0.214, -0.631)]
    scale = 1 + 8.15e-6
    X = tx + scale * xx - rz * yy + ry * zz
    Y = ty + rz * xx + scale * yy - rx * zz
    Z = tz - ry * xx + rx * yy + scale * zz

    # WGS84 cartesian -> geodetic.
    aw = 6378137.0
    ew2 = 6.6943799901413165e-3
    out_lon = math.atan2(Y, X)
    p = math.hypot(X, Y)
    out_lat = math.atan2(Z, p * (1 - ew2))
    for _ in range(8):
        nw = aw / math.sqrt(1 - ew2 * math.sin(out_lat) ** 2)
        out_lat = math.atan2(Z + ew2 * nw * math.sin(out_lat), p)
    return round(math.degrees(out_lon), 6), round(math.degrees(out_lat), 6)


def number(value: object) -> float | None:
    try:
        result = float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def parse_frequency(text: str, fallback: int | None = None) -> int | None:
    match = re.search(r"\b(700|800|900|1800|2100|2300|2600|3600)\s*MHz\b", text, re.I)
    return int(match.group(1)) if match else fallback


def cell_row(headers: list[object], values: tuple[object, ...]) -> dict[str, object]:
    return {str(name or "").strip(): value for name, value in zip(headers, values)}


def build(input_dir: Path, output_file: Path) -> dict[str, object]:
    records: dict[tuple[str, str, int], dict[str, object]] = {}
    dropped = Counter()
    for filename, network in FILES.items():
        workbook_path = input_dir / filename
        if not workbook_path.exists():
            raise FileNotFoundError(f"Missing source workbook: {workbook_path}")
        ws = load_workbook(workbook_path, read_only=True, data_only=True)["Sites"]
        rows = ws.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(rows)]
        for raw in rows:
            row = cell_row(headers, raw)
            identity = str(row.get("Site Identity") or "").strip()
            easting = number(row.get("Easting"))
            northing = number(row.get("Northing"))
            eirp = number(row.get("Max EIRP"))
            band = str(row.get("Band Plan") or row.get("Band") or "").strip()
            if not band and "36ghz" in filename:
                band = "3600MHz LTE"
            frequency = parse_frequency(band, 3600 if "36ghz" in filename else None)
            service = str(row.get("Siteviewer Services") or row.get("Services") or "").upper()
            if not identity or easting is None or northing is None or eirp is None or eirp <= 0 or not frequency:
                dropped["invalid_or_zero_power"] += 1
                continue
            if not (0 <= easting <= 400000 and 0 <= northing <= 470000):
                dropped["outside_irish_grid_extent"] += 1
                continue
            lte = "LTE" in band.upper()
            nr = bool(re.search(r"\bNR(?:5G)?\b", service))
            if not lte and not nr:
                dropped["no_lte_or_nr_flag"] += 1
                continue
            key = (network, identity, frequency)
            current = records.get(key)
            lon, lat = irish_grid_to_wgs84(easting, northing)
            if current is None:
                current = {"network": network, "id": identity, "lon": lon, "lat": lat,
                           "freq": frequency, "eirp": eirp, "lte": lte, "nr": nr}
                records[key] = current
            else:
                if (abs(current["lon"] - lon) > 0.00001 or abs(current["lat"] - lat) > 0.00001):
                    dropped["coordinate_conflicts"] += 1
                current["eirp"] = max(current["eirp"], eirp)
                current["lte"] = current["lte"] or lte
                current["nr"] = current["nr"] or nr

    ordered = sorted(records.values(), key=lambda item: (NETWORKS[item["network"]], item["id"], item["freq"]))
    grouped: dict[tuple[str, str], dict[str, object]] = {}
    counts = Counter()
    for item in ordered:
        network_site = (item["network"], item["id"])
        site = grouped.setdefault(network_site, {
            "network": item["network"], "lon": item["lon"], "lat": item["lat"], "bands": []
        })
        site["bands"].append([item["freq"], item["eirp"], int(item["lte"]) | (int(item["nr"]) << 1)])
        counts[f"{item['network']}_siteBandRows"] += 1
        counts[f"{item['network']}_4gSiteBands"] += bool(item["lte"])
        counts[f"{item['network']}_5gSiteBands"] += bool(item["nr"])
    locations: dict[tuple[str, float, float], dict[str, object]] = {}
    for item in grouped.values():
        location_key = (item["network"], item["lon"], item["lat"])
        location = locations.setdefault(location_key, {
            "network": item["network"], "lon": item["lon"], "lat": item["lat"], "bands": {}
        })
        for frequency, eirp, flags in item["bands"]:
            band = location["bands"].setdefault(frequency, [frequency, eirp, flags])
            band[1] = max(band[1], eirp)
            band[2] |= flags
    sites = []
    for item in sorted(locations.values(), key=lambda value: (NETWORKS[value["network"]], value["lon"], value["lat"])):
        bands = sorted(item["bands"].values(), key=lambda band: band[0])
        sites.append([NETWORKS[item["network"]], item["lon"], item["lat"], bands])
        counts[f"{item['network']}_uniqueSites"] += 1
    payload = {
        "source": "ComReg 2026 Q1 non-confidential mobile licence site schedules",
        "sourceUrl": SOURCE_URL,
        "coordinateReference": "TM75 / Irish Grid (EPSG:29903) transformed to WGS84",
        "tupleFields": ["networkIndex", "longitude", "latitude", "bands"],
        "bandTupleFields": ["frequencyMHz", "maxEirpDbm", "technologyBits"],
        "technologyBits": {"4g": 1, "5g": 2},
        "networks": ["Eir", "Three", "Vodafone"],
        "counts": dict(counts),
        "records": sites,
    }
    output_file.parent.mkdir(parents=True, exist_ok=True)
    output_file.write_text(json.dumps(payload, separators=(",", ":")), encoding="utf-8")
    return {"records": len(sites), "bytes": output_file.stat().st_size,
            "counts": dict(counts), "dropped": dict(dropped)}


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("usage: build-mobile-sites.py <source-xlsx-dir> <output-json>")
    print(json.dumps(build(Path(sys.argv[1]), Path(sys.argv[2])), indent=2))
